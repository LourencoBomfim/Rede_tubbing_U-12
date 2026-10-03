"""Converte controle_tubing_U12.xlsx em data.json (usado pelo site)."""
import json, os
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(HERE, "controle_tubing_U12.xlsx")
OUT = os.path.join(HERE, "data.json")

def norm(x): return str(x if x is not None else "").strip().lower()
def st(x, header=False):
    n = norm(x)
    if n == "montado": return 2
    if "andamento" in n: return 1
    if header and (n == "" or "vis" in n): return -1
    return 0
def num(x, default=0):
    try: return float(x) if x not in (None, "") else default
    except ValueError: return default
def rows(ws):
    it = ws.iter_rows(values_only=True)
    head = [str(c).strip() if c is not None else "" for c in next(it)]
    for r in it:
        if any(c not in (None, "") for c in r):
            yield dict(zip(head, r))

old_blank = {}
if os.path.exists(OUT):
    try:
        old_blank = {p["id"]: bool(p.get("blank")) for p in json.load(open(OUT, encoding="utf-8"))["sp"]}
    except Exception:
        pass

wb = load_workbook(XLSX, data_only=True)
pda, pmap = [], {}
for r in rows(wb["PDAs"]):
    n = str(r["PDA"]).strip()
    o = {"n": n, "z": str(r["Zona"] or ""), "el": int(num(r["Elevação"])), "len": num(r["Metragem PDA (m)"]),
         "h": st(r["Status header"], True), "v": max(0, st(r["Status vaso"])), "rows": []}
    pmap[n] = o; pda.append(o)
for r in rows(wb["Tags_PDA"]):
    o = pmap.get(str(r["PDA"]).strip()); tag = str(r["TAG"] or "").strip()
    if not o or not tag: continue
    li = max(1, int(num(r["Linha"], 1))) - 1
    while len(o["rows"]) <= li: o["rows"].append([])
    o["rows"][li].append({"tag": tag, "s": max(0, st(r["Status"])), "m": num(r["Metragem (m)"])})
for o in pda: o["rows"] = [x for x in o["rows"] if x]

sps, smap = [], {}
for r in rows(wb["StandPipes"]):
    i = str(r["Stand Pipe"]).strip(); tag = str(r["TAG"] or "").strip()
    if not i or not tag: continue
    if i not in smap:
        smap[i] = {"id": i, "row": int(num(r["Fileira"], 1)), "f": "B" if str(r["Forno"]).strip().upper().endswith("B") else "A",
                   "l": [], "r": [], "blank": old_blank.get(i, False)}
        sps.append(smap[i])
    m = r["Metragem (m)"]
    x = {"tag": tag, "s": max(0, st(r["Status"])), "m": None if m in (None, "") else num(m)}
    (smap[i]["l"] if norm(r["Lado"]).startswith("e") else smap[i]["r"]).append(x)
for o in sps:
    o["t"] = o["l"] + o["r"]; o["nl"] = len(o["l"]); del o["l"], o["r"]

json.dump({"pda": pda, "sp": sps}, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print(f"OK: {len(pda)} PDAs, {sum(len(x) for o in pda for x in o['rows'])} tags; {len(sps)} stand pipes, {sum(len(o['t']) for o in sps)} tags")
